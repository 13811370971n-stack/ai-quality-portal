"""
Standalone 8D Report API.
AI assistance per section + Word export, independent of Quality Case flow.
"""
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse, FileResponse
from pydantic import BaseModel
from typing import Optional, Dict
from datetime import datetime
import json
import tempfile

from app.core.ai_engine import chat_completion_stream
from app.core.security import get_current_user
from app.models.user import User

router = APIRouter()

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


SECTION_GUIDE = {
    "d0": "D0 \u7d27\u6025\u54cd\u5e94\u63aa\u65bd (ERA)\uff1a\u95ee\u9898\u53d1\u751f\u540e\u7acb\u5373\u91c7\u53d6\u7684\u884c\u52a8\uff0c\u9632\u6b62\u95ee\u9898\u6270\u52a8\u5ba2\u6237",
    "d1": "D1 \u56e2\u961f\u7ec4\u5efa\uff1a\u8df3\u8dc3\u56e2\u961f\u6210\u5458\u3001\u804c\u8d23\u3001\u56e2\u961f\u957f",
    "d2": "D2 \u95ee\u9898\u63cf\u8ff0\uff1a\u75285W2H\u63cf\u8ff0\u95ee\u9898 - What/When/Where/Who/Which/How/How many",
    "d3": "D3 \u4e34\u65f6\u904f\u5236\u63aa\u65bd (ICA)\uff1a\u9694\u79bb\u3001\u5168\u6570\u68c0\u67e5\u3001\u8fd4\u5de5\u7b49\u63aa\u65bd\uff0c\u5305\u542b\u6709\u6548\u6027\u9a8c\u8bc1",
    "d4": "D4 \u6839\u672c\u539f\u56e0\u5206\u6790\uff1a\u533a\u5206\u53d1\u751f\u539f\u56e0\u548c\u6d41\u51fa\u539f\u56e0\uff0c\u75285Why\u6216\u9c7c\u9aa8\u56fe\u9a8c\u8bc1",
    "d5": "D5 \u6c38\u4e45\u7ea0\u6b63\u63aa\u65bd (PCA)\uff1a\u9488\u5bf9\u6bcf\u4e2a\u6839\u56e0\u7684\u6d88\u9664\u63aa\u65bd",
    "d6": "D6 \u63aa\u65bd\u9a8c\u8bc1\uff1a\u6539\u5584\u524d\u540e\u6570\u636e\u5bf9\u6bd4\uff0c\u8bc1\u660e\u63aa\u65bd\u6709\u6548",
    "d7": "D7 \u9884\u9632\u518d\u53d1\u751f\uff1a\u66f4\u65b0FMEA\u3001\u63a7\u5236\u8ba1\u5212\u3001SOP\u3001\u57f9\u8bad\uff0c\u6a2a\u5411\u5c55\u5f00",
    "d8": "D8 \u56e2\u961f\u8ba4\u53ef\uff1a\u603b\u7ed3\u7ecf\u9a8c\u3001\u5173\u95ed\u786e\u8ba4\u3001\u56e2\u961f\u8868\u5f70",
}


class SectionAssistRequest(BaseModel):
    section: str  # d0..d8
    problem_summary: str
    current_content: Optional[str] = ""
    other_sections: Dict[str, str] = {}
    mode: str = "do"  # "do" | "teach"


@router.get("/guide")
async def get_guide():
    """Get 8D section guidance."""
    return {"sections": SECTION_GUIDE}


@router.post("/assist")
async def assist_section(req: SectionAssistRequest):
    """AI assistance for a specific 8D section (streaming)."""
    guide = SECTION_GUIDE.get(req.section.lower(), "")

    base = """\u4f60\u662f8D\u62a5\u544a\u4e13\u5bb6\uff0c\u719f\u6089AIAG\u548cVDA\u6807\u51c6\u30028D\u662f\u7ed3\u6784\u5316\u95ee\u9898\u89e3\u51b3\u65b9\u6cd5\u3002

\u5f53\u524d\u5e2e\u52a9\u7528\u6237\u5b8c\u6210\uff1a""" + guide + """

\u8981\u6c42\uff1a
- \u5185\u5bb9\u5177\u4f53\u53ef\u6267\u884c\uff0c\u4e0d\u8981\u7a7a\u8bdd
- \u7b26\u5408\u6c7d\u8f66/\u5236\u9020\u4e1a8D\u62a5\u544a\u4e13\u4e1a\u8981\u6c42
- \u4f7f\u7528\u4e2d\u6587"""

    if req.mode == "teach":
        base += """

## \u6559\u5b66\u6a21\u5f0f
\u4e0d\u8981\u76f4\u63a5\u7ed9\u51fa\u5b8c\u6574\u5185\u5bb9\u3002\u800c\u662f\uff1a
1. \u89e3\u91ca\u8fd9\u4e00\u8282\u7684\u8981\u6c42\u548c\u5e38\u89c1\u9519\u8bef
2. \u63d0\u95ee\u5f15\u5bfc\u7528\u6237\u81ea\u5df1\u586b\u5199
3. \u7ed9\u51fa1\u4e2a\u53c2\u8003\u793a\u4f8b\uff08\u4f46\u4e0d\u662f\u7528\u6237\u7684\u5177\u4f53\u5185\u5bb9\uff09"""
    else:
        base += """

## \u5e2e\u4f60\u505a\u6a21\u5f0f
\u76f4\u63a5\u7ed9\u51fa\u53ef\u4ee5\u76f4\u63a5\u586b\u5165\u62a5\u544a\u7684\u5185\u5bb9\u3002\u5982\u679c\u4fe1\u606f\u4e0d\u8db3\uff0c\u57fa\u4e8e\u5408\u7406\u5047\u8bbe\u7ed9\u51fa\uff0c\u5e76\u6807\u6ce8\u9700\u8981\u6838\u5b9e\u7684\u5730\u65b9\u3002"""

    user_content = "\u95ee\u9898\u80cc\u666f: " + req.problem_summary
    if req.current_content:
        user_content += "\n\n\u7528\u6237\u5f53\u524d\u5185\u5bb9:\n" + req.current_content
    if req.other_sections:
        filled = {k: v[:300] for k, v in req.other_sections.items() if v}
        if filled:
            user_content += "\n\n\u5df2\u586b\u5199\u7684\u5176\u4ed6\u90e8\u5206:\n" + json.dumps(filled, ensure_ascii=False)

    messages = [
        {"role": "system", "content": base},
        {"role": "user", "content": user_content},
    ]

    async def generate():
        try:
            async for chunk in chat_completion_stream(messages, temperature=0.6):
                yield "data: " + json.dumps({"content": chunk}) + "\n\n"
        except Exception as e:
            yield "data: " + json.dumps({"error": str(e)}) + "\n\n"
        yield "data: [DONE]\n\n"

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


class ReviewRequest(BaseModel):
    sections: Dict[str, str]
    problem_summary: str = ""


@router.post("/review")
async def review_8d(req: ReviewRequest):
    """AI reviews the complete 8D for logical consistency (streaming)."""
    system = """\u4f60\u662f8D\u5ba1\u6838\u4e13\u5bb6\uff08\u5ba2\u6237\u8d28\u91cf\u4ee3\u8868\u89c6\u89d2\uff09\u3002\u5ba1\u6838\u7528\u6237\u63d0\u4ea4\u76848D\u62a5\u544a\u3002

\u5ba1\u6838\u91cd\u70b9\uff1a
1. **\u903b\u8f91\u4e00\u81f4\u6027**\uff1aD4\u6839\u56e0\u662f\u5426\u771f\u6b63\u89e3\u91caD2\u95ee\u9898\uff1fD5\u63aa\u65bd\u662f\u5426\u9488\u5bf9D4\u6839\u56e0\uff1f
2. **\u5b8c\u6574\u6027**\uff1a\u662f\u5426\u6709\u7f3a\u5931\u7684\u5173\u952e\u4fe1\u606f
3. **\u53ef\u9a8c\u8bc1\u6027**\uff1aD6\u662f\u5426\u6709\u6570\u636e\u652f\u6301
4. **\u9884\u9632\u6709\u6548\u6027**\uff1aD7\u662f\u5426\u80fd\u771f\u6b63\u9632\u6b62\u518d\u53d1
5. **\u5e38\u89c1\u95ee\u9898**\uff1a\u6839\u56e0\u592a\u6d45\uff08\u505c\u5728\u73b0\u8c61\uff09\u3001\u63aa\u65bd\u53ea\u80fd\u53d1\u73b0\u4e0d\u80fd\u6d88\u9664\u3001\u7f3a\u5c11\u6d41\u51fa\u539f\u56e0\u5206\u6790

\u8f93\u51fa\u683c\u5f0f\uff1a
**\u3010\u6574\u4f53\u8bc4\u5206\u3011** X/10

**\u3010\u5173\u952e\u95ee\u9898\u3011**
| \u8282 | \u95ee\u9898 | \u4e25\u91cd\u5ea6 | \u5efa\u8bae |

**\u3010\u5f3a\u9879\u3011**

**\u3010\u5fc5\u987b\u4fee\u6b63\u3011**

\u4e25\u683c\u4f46\u5efa\u8bbe\u6027\u3002\u4f7f\u7528\u4e2d\u6587\u3002"""

    user_content = "8D\u62a5\u544a\u5185\u5bb9:\n" + json.dumps(req.sections, ensure_ascii=False)[:4000]
    if req.problem_summary:
        user_content = "\u95ee\u9898\u80cc\u666f: " + req.problem_summary + "\n\n" + user_content

    messages = [
        {"role": "system", "content": system},
        {"role": "user", "content": user_content},
    ]

    async def generate():
        try:
            async for chunk in chat_completion_stream(messages, temperature=0.4):
                yield "data: " + json.dumps({"content": chunk}) + "\n\n"
        except Exception as e:
            yield "data: " + json.dumps({"error": str(e)}) + "\n\n"
        yield "data: [DONE]\n\n"

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


class ExportRequest(BaseModel):
    title: str = "8D Report"
    report_no: str = ""
    customer: str = ""
    part_number: str = ""
    owner: str = ""
    sections: Dict[str, str] = {}


def _shade(cell, color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): color, qn("w:val"): "clear"})
    tcPr.append(shd)


@router.post("/export")
async def export_8d(req: ExportRequest):
    """Export standalone 8D to Word."""
    if not HAS_DOCX:
        raise HTTPException(status_code=500, detail="python-docx not installed")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run("8D PROBLEM SOLVING REPORT")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(5, 28, 44)

    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r2 = sub.add_run("AI Quality Portal | aidmaic.top")
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(139, 157, 175)

    doc.add_paragraph()

    # Info table
    t = doc.add_table(rows=3, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    info = [
        ("Report No.", req.report_no or "8D-" + datetime.now().strftime("%Y%m%d"), "Date", datetime.now().strftime("%Y-%m-%d")),
        ("Title", req.title or "-", "Owner", req.owner or "-"),
        ("Customer", req.customer or "-", "Part No.", req.part_number or "-"),
    ]
    for i, row_data in enumerate(info):
        for j, val in enumerate(row_data):
            cell = t.cell(i, j)
            cell.text = str(val)
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(3)
                para.paragraph_format.space_after = Pt(3)
                for run in para.runs:
                    run.font.size = Pt(9)
                    if j % 2 == 0:
                        run.bold = True
            if j % 2 == 0:
                _shade(cell, "F5F7FA")

    doc.add_paragraph()

    # D sections
    titles = {
        "d0": "D0 - Emergency Response Action (ERA)",
        "d1": "D1 - Team Formation",
        "d2": "D2 - Problem Description",
        "d3": "D3 - Interim Containment Action (ICA)",
        "d4": "D4 - Root Cause Analysis",
        "d5": "D5 - Permanent Corrective Action (PCA)",
        "d6": "D6 - Verification of Effectiveness",
        "d7": "D7 - Prevention of Recurrence",
        "d8": "D8 - Team Recognition & Closure",
    }

    for key in ["d0", "d1", "d2", "d3", "d4", "d5", "d6", "d7", "d8"]:
        hp = doc.add_paragraph()
        hr = hp.add_run(titles[key])
        hr.bold = True
        hr.font.size = Pt(12)
        hr.font.color.rgb = RGBColor(5, 28, 44)
        hp.paragraph_format.space_before = Pt(10)
        hp.paragraph_format.space_after = Pt(4)

        content = req.sections.get(key, "").strip()
        if content:
            for line in content.split("\n"):
                cp = doc.add_paragraph(line.strip())
                cp.paragraph_format.left_indent = Cm(0.5)
                cp.paragraph_format.space_after = Pt(3)
                for run in cp.runs:
                    run.font.size = Pt(10)
        else:
            cp = doc.add_paragraph("(Not completed)")
            cp.paragraph_format.left_indent = Cm(0.5)
            if cp.runs:
                cp.runs[0].font.size = Pt(10)
                cp.runs[0].font.color.rgb = RGBColor(139, 157, 175)

    # Footer
    doc.add_paragraph()
    fp = doc.add_paragraph()
    fp.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    fr = fp.add_run("Generated: " + datetime.now().strftime("%Y-%m-%d %H:%M") + " | AI Quality Portal")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(139, 157, 175)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(tmp.name)
    tmp.close()

    filename = "8D_" + (req.report_no or datetime.now().strftime("%Y%m%d")) + ".docx"
    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=filename,
    )
