"""
File upload with text extraction for AI context.
Stores files locally, extracts text content into CaseEvidence.content.
"""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.orm import Session
import os
import io
import uuid
from datetime import datetime, timezone

from app.database.session import get_db
from app.models.user import User
from app.models.quality_case import QualityCase
from app.models.evidence import CaseEvidence, CaseTimeline
from app.core.security import require_user

router = APIRouter()

UPLOAD_DIR = "/root/Projects/ai-quality-portal/backend/uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

ALLOWED_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".csv", ".txt", ".md",
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp",
    ".ppt", ".pptx",
}

MAX_EXTRACT_CHARS = 8000


def get_extension(filename: str) -> str:
    return os.path.splitext(filename)[1].lower()


def get_evidence_type(ext: str) -> str:
    if ext in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}:
        return "image"
    if ext in {".xls", ".xlsx", ".csv"}:
        return "excel"
    if ext in {".ppt", ".pptx"}:
        return "presentation"
    return "document"


def extract_text(content: bytes, ext: str, filename: str) -> tuple:
    """
    Extract text from file bytes.
    Returns (extracted_text, extraction_note).
    """
    try:
        if ext in (".txt", ".md", ".csv"):
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError:
                text = content.decode("gbk", errors="replace")
            if ext == ".csv":
                import pandas as pd
                df = pd.read_csv(io.BytesIO(content))
                summary = "CSV: " + str(len(df)) + " rows x " + str(len(df.columns)) + " cols\n"
                summary += "Columns: " + ", ".join(str(c) for c in df.columns) + "\n\n"
                summary += df.head(20).to_string()
                if len(df) > 20:
                    summary += "\n... (" + str(len(df) - 20) + " more rows)\n\n"
                num = df.select_dtypes(include="number")
                if not num.empty:
                    summary += "\nNumeric summary:\n" + num.describe().to_string()
                return summary[:MAX_EXTRACT_CHARS], "csv_parsed"
            return text[:MAX_EXTRACT_CHARS], "text_read"

        if ext in (".xlsx", ".xls"):
            import pandas as pd
            sheets = pd.read_excel(io.BytesIO(content), sheet_name=None)
            parts = []
            for name, df in list(sheets.items())[:5]:
                parts.append("=== Sheet: " + str(name) + " (" + str(len(df)) + " rows x " + str(len(df.columns)) + " cols) ===")
                parts.append("Columns: " + ", ".join(str(c) for c in df.columns))
                parts.append(df.head(15).to_string())
                num = df.select_dtypes(include="number")
                if not num.empty:
                    parts.append("Numeric summary:\n" + num.describe().to_string())
                parts.append("")
            return "\n".join(parts)[:MAX_EXTRACT_CHARS], "excel_parsed"

        if ext == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables[:10]:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            return "\n".join(parts)[:MAX_EXTRACT_CHARS], "docx_parsed"

        if ext == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                try:
                    from PyPDF2 import PdfReader
                except ImportError:
                    return "", "pdf_lib_missing"
            reader = PdfReader(io.BytesIO(content))
            parts = []
            for i, page in enumerate(reader.pages[:30]):
                try:
                    t = page.extract_text()
                    if t and t.strip():
                        parts.append("--- Page " + str(i + 1) + " ---\n" + t.strip())
                except Exception:
                    continue
            if not parts:
                return "", "pdf_no_text_layer"
            return "\n\n".join(parts)[:MAX_EXTRACT_CHARS], "pdf_parsed"

        if ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"):
            return "", "image_no_ocr"

        return "", "unsupported_for_extraction"

    except Exception as e:
        return "", "extract_error: " + str(e)[:100]


@router.post("/cases/{case_id}/upload")
async def upload_file(
    case_id: int,
    file: UploadFile = File(...),
    description: str = Form(default=""),
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    """Upload a file as evidence, extracting text content for AI context."""
    case = db.query(QualityCase).filter(
        QualityCase.id == case_id, QualityCase.user_id == user.id
    ).first()
    if not case:
        raise HTTPException(status_code=404, detail="Case not found")

    ext = get_extension(file.filename or "")
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="File type " + ext + " not supported")

    file_id = str(uuid.uuid4())[:8]
    safe_name = str(case_id) + "_" + file_id + ext
    file_path = os.path.join(UPLOAD_DIR, safe_name)

    content = await file.read()
    with open(file_path, "wb") as f:
        f.write(content)

    # Extract text for AI context
    extracted, note = extract_text(content, ext, file.filename or safe_name)

    stored_content = description or ""
    if extracted:
        if stored_content:
            stored_content += "\n\n"
        stored_content += "[\u63d0\u53d6\u5185\u5bb9]\n" + extracted
    elif not stored_content:
        stored_content = "\u6587\u4ef6: " + (file.filename or safe_name) + " (" + note + ")"

    evidence = CaseEvidence(
        case_id=case_id,
        evidence_type=get_evidence_type(ext),
        source="user_provided",
        title=file.filename or safe_name,
        content=stored_content,
        file_path=file_path,
        verification_status="unverified",
        related_step=case.current_step,
        created_by="user",
    )
    db.add(evidence)
    db.add(CaseTimeline(
        case_id=case_id, event_type="evidence_added",
        description="\u4e0a\u4f20\u6587\u4ef6: " + (file.filename or safe_name) + (" (\u5df2\u63d0\u53d6\u6587\u672c)" if extracted else ""),
        actor="user",
    ))
    db.commit()
    db.refresh(evidence)

    return {
        "id": evidence.id,
        "title": evidence.title,
        "evidence_type": evidence.evidence_type,
        "file_path": safe_name,
        "extracted": bool(extracted),
        "extract_note": note,
        "extract_chars": len(extracted),
        "preview": extracted[:300] if extracted else None,
        "created_at": evidence.created_at.isoformat() if evidence.created_at else None,
    }


@router.get("/cases/{case_id}/evidences")
async def list_evidences(case_id: int, user: User = Depends(require_user), db: Session = Depends(get_db)):
    """List all evidences for a case, with content preview."""
    case = db.query(QualityCase).filter(
        QualityCase.id == case_id, QualityCase.user_id == user.id
    ).first()
    if not case:
        raise HTTPException(status_code=404, detail="Case not found")

    evidences = db.query(CaseEvidence).filter(CaseEvidence.case_id == case_id).all()
    return [
        {
            "id": e.id,
            "evidence_type": e.evidence_type,
            "source": e.source,
            "title": e.title,
            "content_preview": (e.content[:400] if e.content else None),
            "has_extracted_text": bool(e.content and "[\u63d0\u53d6\u5185\u5bb9]" in e.content),
            "confidence": e.confidence,
            "verification_status": e.verification_status,
            "related_step": e.related_step,
            "created_at": e.created_at.isoformat() if e.created_at else None,
            "created_by": e.created_by,
        }
        for e in evidences
    ]


@router.delete("/cases/{case_id}/evidences/{evidence_id}")
async def delete_evidence(case_id: int, evidence_id: int,
                          user: User = Depends(require_user), db: Session = Depends(get_db)):
    """Delete an evidence item and its file."""
    case = db.query(QualityCase).filter(
        QualityCase.id == case_id, QualityCase.user_id == user.id
    ).first()
    if not case:
        raise HTTPException(status_code=404, detail="Case not found")

    ev = db.query(CaseEvidence).filter(
        CaseEvidence.id == evidence_id, CaseEvidence.case_id == case_id
    ).first()
    if not ev:
        raise HTTPException(status_code=404, detail="Evidence not found")

    if ev.file_path and os.path.exists(ev.file_path):
        try:
            os.remove(ev.file_path)
        except OSError:
            pass

    db.delete(ev)
    db.commit()
    return {"status": "deleted"}
